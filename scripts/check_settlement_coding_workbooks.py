"""Independent reader checks for nonempty, synthetic API-produced XLSX exports."""
import hashlib
import json
from pathlib import Path
import sys
from zipfile import ZipFile

import openpyxl


output = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent / "outputs"
expected_sheets = ["月度概览", "讲师排名", "机构汇总", "对账明细", "说明"]
checks = 0
results = []


def check(condition, message):
    global checks
    checks += 1
    if not condition:
        raise AssertionError(message)


for mode, count in (("teaching", 2), ("payment", 3)):
    path = output / f"actual-coded-{mode}-synthetic.xlsx"
    workbook = openpyxl.load_workbook(path, data_only=False)
    check(workbook.sheetnames == expected_sheets, "exact five approved sheets")
    detail = workbook["对账明细"]
    check(detail.max_row - 3 == count, "nonempty detail count")
    amounts = [detail.cell(row, 13).value for row in range(4, detail.max_row + 1)]
    check(sum(amounts) == 80, "detail net amount is 80")
    check(sorted(amounts) == ([-70, 150] if mode == "teaching" else [-70, 50, 100]), "signed corrections/payments/refund preserved")
    check(workbook["月度概览"]["B4"].value == 80, "monthly total agrees")
    check(workbook["讲师排名"]["G4"].value == 80, "teacher total agrees")
    check(workbook["机构汇总"]["F4"].value == 80, "organization total agrees")
    for row in range(4, detail.max_row + 1):
        for col, code in ((6, "001"), (7, "00020"), (8, "00030")):
            cell = detail.cell(row, col)
            check(cell.value == code and cell.data_type == "s", "leading-zero code retained as text")
        if mode == "payment":
            check(all(detail.cell(row, col).value == "不归集" for col in (9, 10, 11, 12, 15, 16, 17, 18)), "payment/refund never duplicates hours")
    if mode == "teaching":
        check(all(workbook["月度概览"].cell(row, 2).value == 2 for row in range(6, 10)), "all four hour summaries use latest two-hour state")
    else:
        check(all(workbook["月度概览"].cell(row, 2).value == "不归集" for row in range(6, 10)), "cash overview has no teaching hours")
    notes = list(workbook["说明"].values)
    coding_note = [row for row in notes if row[0] == "编码来源"]
    check(len(coding_note) == 1 and "审计关联" in coding_note[0][1], "coding provenance explained")
    for sheet in workbook:
        check(sheet.freeze_panes == "A4", "three header rows remain frozen")
        check(all(cell.data_type != "f" for row in sheet for cell in row), "no formulas")
    with ZipFile(path) as archive:
        text = "\n".join(archive.read(name).decode("utf-8") for name in archive.namelist())
        check("TargetMode=\"External\"" not in text, "no external references")
        check(not any(s in text for s in ("PRIVATE-", "catalog_scope_id", "teacher_id", "evidence_note")), "no private fixture values or internal identity fields")
    results.append({"file": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "detail_rows": count, "net_amount": "80.00", "signed_amounts": amounts, "code_cells_are_text": True})

evidence = {"reader": "openpyxl", "reader_version": openpyxl.__version__, "source_kind": "REAL_API_HANDLE_SYNTHETIC_XLSX", "checks": checks, "workbooks": results}
(output / "workbook-evidence.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n")
print(json.dumps(evidence, ensure_ascii=False, indent=2))
